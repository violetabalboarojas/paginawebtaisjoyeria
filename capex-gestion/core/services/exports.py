"""Exportación a Excel (openpyxl) y PDF (reportlab)."""

import datetime
import io
from xml.sax.saxutils import escape
from decimal import Decimal

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from ..models import CURRENCY_SYMBOL, SystemSettings

TEAL = colors.HexColor("#1F7F78")
ORANGE = colors.HexColor("#CF7D3A")
LIGHT = colors.HexColor("#EEF6F5")
XL_HEAD = PatternFill("solid", fgColor="1F7F78")
MONEY_FMT = '#,##0.00'


def _fmt(v, typ=None):
    if v is None:
        return ""
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, Decimal):
        return f"{v:,.2f}" + ("%" if typ == "pct" else "")
    return str(v)


def _xl_response(wb, name):
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{name}.xlsx"'
    return resp


def _pdf_response(buf, name):
    resp = HttpResponse(buf.getvalue(), content_type="application/pdf")
    resp["Content-Disposition"] = f'attachment; filename="{name}.pdf"'
    return resp


def _write_table(ws, start_row, headers, rows, types):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = XL_HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r, row in enumerate(rows, start_row + 1):
        for c, (v, t) in enumerate(zip(row, types), 1):
            if isinstance(v, Decimal):
                v = float(v)
            cell = ws.cell(row=r, column=c, value=v)
            if t == "money":
                cell.number_format = MONEY_FMT
            elif t == "date" and v:
                cell.number_format = "DD/MM/YYYY"
    for c, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(c)].width = max(12, min(40, len(str(h)) + 4))
    return start_row + len(rows) + 1


def _styles():
    ss = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("t", parent=ss["Title"], fontName="Helvetica-Bold", fontSize=15, textColor=TEAL, alignment=0, spaceAfter=2),
        "sub": ParagraphStyle("s", parent=ss["Normal"], fontSize=8.5, textColor=colors.HexColor("#5F6B6D")),
        "h": ParagraphStyle("h", parent=ss["Heading3"], fontSize=10.5, textColor=ORANGE, spaceBefore=8, spaceAfter=4),
        "cell": ParagraphStyle("c", parent=ss["Normal"], fontSize=7.5, leading=9),
        "small": ParagraphStyle("x", parent=ss["Normal"], fontSize=7.5, leading=9.5, textColor=colors.HexColor("#3A4446")),
    }


def _pdf_table(data, col_widths=None, money_cols=(), header=True):
    t = Table(data, colWidths=col_widths, repeatRows=1 if header else 0)
    style = [
        ("FONT", (0, 0), (-1, -1), "Helvetica", 7.5),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D5DEDD")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]),
    ]
    if header:
        style += [("BACKGROUND", (0, 0), (-1, 0), TEAL), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                  ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 7.5)]
    for c in money_cols:
        style.append(("ALIGN", (c, 1 if header else 0), (c, -1), "RIGHT"))
    t.setStyle(TableStyle(style))
    return t


def _footer(canvas, doc):
    cfg = SystemSettings.load()
    canvas.saveState()
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#7A8587"))
    canvas.drawString(12 * mm, 8 * mm, f"{cfg.company_name} · Generado el {timezone.localtime():%d/%m/%Y %H:%M}")
    canvas.drawRightString(doc.pagesize[0] - 12 * mm, 8 * mm, f"Página {doc.page}")
    canvas.restoreState()


# ------------------------------------------------------------------------------------ reportes
def report_xlsx(report, filters_text):
    wb = Workbook()
    ws = wb.active
    ws.title = report.title[:30]
    cfg = SystemSettings.load()
    ws["A1"] = f"{cfg.company_name} · {report.title}"
    ws["A1"].font = Font(bold=True, size=13, color="1F7F78")
    ws["A2"] = filters_text
    headers = [c[1] for c in report.columns]
    types = [c[2] for c in report.columns]
    rows = [[r.get(k) for k, _, _ in report.columns] for r in report.rows]
    end = _write_table(ws, 4, headers, rows, types)
    for cur, tot in report.totals.items():
        end += 1
        ws.cell(row=end, column=1, value=f"Total {cur}").font = Font(bold=True)
        for c, (k, _, t) in enumerate(report.columns, 1):
            if t == "money":
                cell = ws.cell(row=end, column=c, value=float(tot.get(k, 0)))
                cell.number_format = MONEY_FMT
                cell.font = Font(bold=True)
    if report.notes:
        ws.cell(row=end + 2, column=1, value=report.notes)
    return _xl_response(wb, f"reporte_{report.title.lower().replace(' ', '_')}")


def report_pdf(report, filters_text):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=12 * mm, bottomMargin=14 * mm)
    st = _styles()
    cfg = SystemSettings.load()
    story = [Paragraph(report.title, st["title"]), Paragraph(escape(f"{cfg.company_name} · {filters_text}"), st["sub"]), Spacer(1, 6)]
    head = [c[1] for c in report.columns]
    data = [[Paragraph(escape(h), st["cell"]) for h in head]]
    for r in report.rows:
        data.append([Paragraph(escape(_fmt(r.get(k), t)), st["cell"]) for k, _, t in report.columns])
    for cur, tot in report.totals.items():
        data.append([Paragraph(f"<b>Total {cur}</b>", st["cell"])] +
                    [Paragraph(f"<b>{_fmt(tot.get(k))}</b>", st["cell"]) if t == "money" else "" for k, _, t in report.columns[1:]])
    money_cols = [i for i, c in enumerate(report.columns) if c[2] in ("money", "pct", "int")]
    if len(data) == 1:
        story.append(Paragraph("Sin resultados para los filtros seleccionados.", st["small"]))
    else:
        story.append(_pdf_table(data, money_cols=money_cols))
    if report.notes:
        story += [Spacer(1, 6), Paragraph(escape(report.notes), st["small"])]
    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return _pdf_response(buf, f"reporte_{report.title.lower().replace(' ', '_')}")


# ------------------------------------------------------------------------------------ estado de cuenta
def _statement_blocks(ctx):
    op, s, l = ctx["op"], ctx["s"], ctx["ltv"]
    sym = CURRENCY_SYMBOL[op.currency]
    m = lambda v: f"{sym} {v:,.2f}" if v is not None else "–"  # noqa: E731
    general = [
        ("Operación", op.code), ("Estado", op.get_status_display()), ("Cliente", f"{op.client} ({op.client.doc_type} {op.client.doc_number})"),
        ("Inversionista", str(op.investor)), ("Capital inicial", m(op.principal)), ("Capital pendiente", m(s["capital_pending"])),
        ("Tasa", f"{op.get_rate_type_display()} {op.rate_value.normalize():f}% · {op.get_interest_method_display()} · base {op.day_base}"),
        ("Periodicidad / amortización", f"{op.get_periodicity_display()} · {op.get_amortization_display()}"),
        ("Fecha inicial", op.start_date.strftime("%d/%m/%Y")), ("Fecha de vencimiento", op.maturity_date.strftime("%d/%m/%Y")),
        ("Intereses generados (devengados)", m(s["interest_accrued"])), ("Intereses pagados", m(s["interest_paid"])),
        ("Intereses pendientes", m(s["interest_pending"])), ("Penalidades generadas", m(s["penalty_generated"])),
        ("Penalidades pagadas / condonadas", f"{m(s['penalty_paid'])} / {m(s['penalty_waived'])}"),
        ("Penalidades pendientes", m(s["penalty_pending"])), ("Total pagado", m(s["total_paid"])),
        ("SALDO TOTAL AL CORTE", m(s["total_due"])), ("Días de atraso", str(s["days_late"])),
        ("Valor de la garantía", m(l["guarantee_value"]) if not l["missing"] else "Sin garantía registrada"),
        ("LTV (capital / valor garantía)", f"{l['ltv']}% (límite {l['limit']}%)" if l["ltv"] is not None else "–"),
    ]
    sched_head = ["N°", "Vencimiento", "Capital", "Interés", "Penalidad", "Total esperado", "Pagado", "Saldo", "Fecha pago", "Días atraso", "Estado"]
    sched = [[r.number, r.due_date, r.capital, r.interest, r.penalty_amount, r.expected_total, r.paid_total, r.balance,
              r.payment_date, r.days_late, r.get_status_display()] for r in ctx["rows"]]
    pay_head = ["Código", "Fecha", "Concepto", "Monto", "Capital", "Interés", "Penalidad", "Medio", "N° operación", "Estado"]
    pays = [[p.code, p.payment_date, p.get_concept_display(), p.amount, p.capital, p.interest, p.penalty, p.get_method_display(),
             p.bank_reference, p.get_status_display()] for p in ctx["payments"]]
    guar = [[m_.property.get_property_type_display(), m_.property.registry_number, m_.property.registry_office,
             f"{m_.property.district}, {m_.property.province}", m_.get_rank_display(), m_.get_inscription_status_display(),
             m_.property.commercial_value, m_.property.realization_value] for m_ in ctx["mortgages"]]
    return general, sched_head, sched, pay_head, pays, guar


def statement_xlsx(ctx):
    op = ctx["op"]
    general, sh, sched, ph, pays, guar = _statement_blocks(ctx)
    wb = Workbook()
    ws = wb.active
    ws.title = "Estado de cuenta"
    ws["A1"] = f"Estado de cuenta {op.code} al {ctx['as_of']:%d/%m/%Y}"
    ws["A1"].font = Font(bold=True, size=13, color="1F7F78")
    for i, (k, v) in enumerate(general, 3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 60
    ws2 = wb.create_sheet("Cronograma")
    _write_table(ws2, 1, sh, sched, ["int", "date"] + ["money"] * 6 + ["date", "int", "text"])
    ws3 = wb.create_sheet("Pagos")
    _write_table(ws3, 1, ph, pays, ["text", "date", "text"] + ["money"] * 4 + ["text"] * 3)
    ws4 = wb.create_sheet("Garantía")
    _write_table(ws4, 1, ["Tipo", "Partida", "Oficina", "Ubicación", "Rango", "Inscripción", "Valor comercial", "Valor realización"],
                 guar, ["text"] * 6 + ["money", "money"])
    return _xl_response(wb, f"estado_cuenta_{op.code}")


def statement_pdf(ctx):
    op = ctx["op"]
    general, sh, sched, ph, pays, guar = _statement_blocks(ctx)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=14 * mm)
    st = _styles()
    cfg = SystemSettings.load()
    story = [Paragraph(f"Estado de cuenta · {op.code}", st["title"]),
             Paragraph(escape(cfg.company_name) + f" · Corte al {ctx['as_of']:%d/%m/%Y}", st["sub"]), Spacer(1, 6)]
    half = (len(general) + 1) // 2
    left, right = general[:half], general[half:]
    grid = []
    for i in range(half):
        a = left[i]
        b = right[i] if i < len(right) else ("", "")
        grid.append([Paragraph(f"<b>{escape(a[0])}</b>", st["cell"]), Paragraph(escape(a[1]), st["cell"]),
                     Paragraph(f"<b>{escape(b[0])}</b>", st["cell"]), Paragraph(escape(b[1]), st["cell"])])
    story.append(_pdf_table(grid, col_widths=[38 * mm, 55 * mm, 38 * mm, 55 * mm], header=False))
    story.append(Paragraph("Cronograma de pagos", st["h"]))
    data = [sh] + [[_fmt(v) for v in row] for row in sched]
    story.append(_pdf_table(data, money_cols=range(2, 8)))
    story.append(Paragraph("Pagos realizados", st["h"]))
    if pays:
        story.append(_pdf_table([ph] + [[_fmt(v) for v in row] for row in pays], money_cols=range(3, 7)))
    else:
        story.append(Paragraph("Sin pagos registrados.", st["small"]))
    story.append(Paragraph("Garantía hipotecaria", st["h"]))
    if guar:
        story.append(_pdf_table([["Tipo", "Partida", "Oficina", "Ubicación", "Rango", "Inscripción", "V. comercial", "V. realización"]] +
                                [[_fmt(v) for v in row] for row in guar], money_cols=(6, 7)))
    else:
        story.append(Paragraph("Sin garantía registrada.", st["small"]))
    story += [Spacer(1, 8), Paragraph("Saldo total al corte = capital pendiente + intereses devengados no pagados + penalidades pendientes. "
                                      "Documento informativo generado por el sistema.", st["small"])]
    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return _pdf_response(buf, f"estado_cuenta_{op.code}")


def audit_xlsx(logs):
    wb = Workbook()
    ws = wb.active
    ws.title = "Auditoría"
    rows = [[l.created_at.astimezone(timezone.get_current_timezone()).strftime("%d/%m/%Y"),
             l.created_at.astimezone(timezone.get_current_timezone()).strftime("%H:%M:%S"), l.username, l.get_action_display(),
             l.module, l.record_label, str(l.old_value or ""), str(l.new_value or ""), l.ip_address or "", l.detail] for l in logs]
    _write_table(ws, 1, ["Fecha", "Hora", "Usuario", "Acción", "Módulo", "Registro", "Valor anterior", "Valor nuevo", "IP", "Detalle"],
                 rows, ["text"] * 10)
    return _xl_response(wb, "auditoria")
