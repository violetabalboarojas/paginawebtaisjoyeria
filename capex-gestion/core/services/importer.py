"""
Importación de cronogramas de bonos titulizados desde Excel.

Reconoce libros con la estructura del "Cronograma PF Fundo La Estrella": bloques por serie
("Monto Emisión", "Valor Nominal", "Cant. Bonos", "Moneda", "Plazo", "TNM", "TEA" + tabla
Periodo/Fecha/Principal/Amortización/Interés/Cuota/Estado), un bloque consolidado con
"Com. Adm." y una hoja de IMC con tasas máximas y la distribución por bonista.

Flujo: parse_workbook() → build_preview() (no escribe nada) → confirm_import() (transacción).
"""

import datetime
import hashlib
import re
import unicodedata
from collections import defaultdict
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

from audit.services import log, snapshot

from ..models import (
    BondHolding, Investor, Operation, Payment, PaymentAllocation, PaymentSchedule, SystemSettings, Trust,
)
from .engine import ONE, ZERO, add_months, compute_penalty, money, unit_interest
from .portfolio import BusinessError, refresh_operation, save_with_code, today

CURRENCY_WORDS = {"dolares": "USD", "dólares": "USD", "usd": "USD", "us$": "USD", "soles": "PEN", "pen": "PEN", "s/": "PEN"}
HEADER_KEYS = {
    "periodo": "period", "fecha": "date", "principal": "principal", "amortizacion": "amort", "interes": "interest",
    "cuota": "installment", "estado": "status", "com. adm.": "fee", "total a pagar": "total",
}


def norm(text):
    text = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", text).strip().lower()


def D(x):
    if x is None or x == "":
        return None
    try:
        return Decimal(str(x))
    except InvalidOperation:
        return None


def as_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    return None


# ------------------------------------------------------------------------------- lector de celdas
class Book:
    """Valor de una celda: el valor guardado por Excel o, si falta, una evaluación de fórmulas simples."""

    REF = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")

    def __init__(self, path):
        self.wf = load_workbook(path, data_only=False)
        self.wv = load_workbook(path, data_only=True)
        self.notes = []

    def raw(self, ws, row, col):
        return self.wf[ws][f"{get_column_letter(col)}{row}"].value

    def value(self, ws, row, col, depth=0):
        cached = self.wv[ws].cell(row, col).value
        if cached is not None:
            return cached
        f = self.wf[ws].cell(row, col).value
        if not isinstance(f, str) or not f.startswith("=") or depth > 30:
            return f if not hasattr(f, "text") else None
        return self._eval(ws, f[1:].lstrip("+"), depth)

    def _cell(self, ws, ref, depth):
        m = self.REF.fullmatch(ref.strip())
        if not m:
            return None
        return self.value(ws, int(m.group(2)), column_index_from_string(m.group(1)), depth + 1)

    def _eval(self, ws, expr, depth):
        expr = expr.strip()
        m = re.fullmatch(r"EDATE\(([^,]+),\s*(-?\d+)\)", expr, re.I)
        if m:
            base = as_date(self._cell(ws, m.group(1), depth))
            return add_months(base, int(m.group(2))) if base else None
        m = re.fullmatch(r"SUM\(([A-Z$]+\d+):([A-Z$]+\d+)\)", expr, re.I)
        if m:
            c1, r1, c2, r2 = range_boundaries(f"{m.group(1)}:{m.group(2)}".replace("$", ""))
            total = Decimal(0)
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    v = D(self.value(ws, r, c, depth + 1))
                    total += v or 0
            return total
        if re.fullmatch(r"[A-Z$0-9+\-*/(). %]+", expr):
            def repl(mm):
                v = self._cell(ws, mm.group(0), depth)
                v = as_date(v).toordinal() if as_date(v) else D(v)
                return f"Decimal('{v if v is not None else 0}')"
            py = self.REF.sub(repl, expr.replace("%", "/100"))
            py = re.sub(r"(?<![\w'.])(\d+(?:\.\d+)?)(?![\w'])", r"Decimal('\1')", py)
            try:
                return eval(py, {"__builtins__": {}}, {"Decimal": Decimal})  # solo números y operadores (validado arriba)
            except Exception:  # noqa: BLE001
                return None
        return None  # fórmula no soportada: se recalcula con el motor


# ------------------------------------------------------------------------------- lectura del libro
def parse_workbook(path):
    book = Book(path)
    out = {"sheets": book.wf.sheetnames, "schedules": [], "trust_name": "", "imc": None, "issues": []}
    for ws in book.wf.worksheets:
        blocks = _parse_blocks(book, ws.title, out)
        if blocks:
            out["schedules"].append({"sheet": ws.title, "blocks": blocks,
                                     "has_prepago": any(r["label"].upper() == "PREPAGO" for b in blocks for r in b["rows"])})
    imc_ws = next((ws for ws in book.wf.worksheets if _find(book, ws.title, lambda t: "int. moratorio" in t or t.startswith("tasa maxima") or t in ("inversionista", "bonista"))), None)
    if imc_ws:
        out["imc"] = _parse_imc(book, imc_ws.title)
    refs = {}
    for name, dn in book.wf.defined_names.items():
        m = re.match(r"'?([^'!]+)'?!", dn.attr_text or "")
        if m:
            refs[name] = m.group(1)
    out["named_ranges"] = refs
    imc_sheets = {refs.get(n) for n in ("serie_a", "serie_b", "consolidado") if refs.get(n)}
    names = [s["sheet"] for s in out["schedules"]]
    rec = next((s for s in names if s in imc_sheets), None) or next((s["sheet"] for s in out["schedules"] if s["has_prepago"]), None)
    out["recommended_sheet"] = rec or (names[-1] if names else None)
    out["issues"] += _compare_sheets(out["schedules"])
    if not out["schedules"]:
        out["issues"].append({"level": "error", "msg": "No se encontró ningún cronograma con la estructura esperada "
                                                        "(celda “Monto Emisión” y tabla Periodo/Fecha/Principal/Amortización/Interés)."})
    return out


def _find(book, ws, pred):
    for row in book.wf[ws].iter_rows():
        for c in row:
            if isinstance(c.value, str) and pred(norm(c.value)):
                return c
    return None


def _parse_blocks(book, ws, out):
    sheet = book.wf[ws]
    blocks = []
    anchors = [c for row in sheet.iter_rows() for c in row if isinstance(c.value, str) and norm(c.value).startswith("monto emision")]
    for a in anchors:
        r, col = a.row, a.column
        params = {}
        for rr in range(r, r + 8):
            label = norm(book.raw(ws, rr, col))
            val = book.value(ws, rr, col + 1)
            if label.startswith("monto emision"):
                params["amount"] = D(val)
            elif label.startswith("valor nominal"):
                params["nominal"] = D(val)
            elif label.startswith("cant"):
                params["units"] = D(val)
            elif label.startswith("moneda"):
                params["currency"] = CURRENCY_WORDS.get(norm(val), None)
            elif label.startswith("plazo"):
                params["term"] = D(val)
            elif label.startswith("tnm"):
                params["tnm"] = D(val)
            elif label.startswith("tea"):
                params["tea"] = D(val)
        title = ""
        for rr in range(r - 1, max(0, r - 5), -1):
            t = book.raw(ws, rr, col)
            if isinstance(t, str) and t.strip():
                if not title and ("serie" in norm(t) or "consolidado" in norm(t) or "bono" in norm(t)):
                    title = t.strip()
                m = re.search(r"patrimonio en fideicomiso\s+(.+?)(\s+d\.?\s?leg|,|$)", norm(t))
                if m and not out["trust_name"]:
                    raw = t.strip()
                    out["trust_name"] = raw[:raw.lower().find(" d.leg")] if " d.leg" in raw.lower() else raw.split(",")[0]
        consolidated = "consolidado" in norm(title)
        # cabecera de la tabla
        hdr_row = next((rr for rr in range(r, r + 15) if norm(book.raw(ws, rr, col)) == "periodo"), None)
        if not hdr_row:
            out["issues"].append({"level": "warning", "msg": f"{ws}: el bloque “{title or a.coordinate}” no tiene tabla de cronograma."})
            continue
        cols = {}
        c = col
        while c < col + 14:
            key = HEADER_KEYS.get(norm(book.raw(ws, hdr_row, c)))
            if key and key not in cols:
                cols[key] = c
            elif book.raw(ws, hdr_row, c) is None and c > col + 6:
                break
            c += 1
        rows = []
        prev_date = None
        rr = hdr_row + 1
        while rr < hdr_row + 80:
            p = book.value(ws, rr, cols.get("period", col))
            if isinstance(p, str) and norm(p).startswith("total"):
                break
            date = as_date(book.value(ws, rr, cols["date"])) if "date" in cols else None
            if date is None:
                if p is None:
                    break
                rr += 1
                continue
            row = {
                "label": str(p if p is not None else "").replace(".0", "").strip(),
                "date": date.isoformat(),
                "principal": _s(book.value(ws, rr, cols["principal"])) if "principal" in cols else None,
                "amort": _s(book.value(ws, rr, cols["amort"])) if "amort" in cols else "0",
                "interest_excel": _s(book.value(ws, rr, cols["interest"])) if "interest" in cols else None,
                "status": str(book.value(ws, rr, cols["status"]) or "").strip() if "status" in cols else None,
                "fee": _s(book.value(ws, rr, cols["fee"])) if "fee" in cols else None,
                "row": rr,
            }
            if prev_date and date < prev_date:
                out["issues"].append({"level": "warning", "msg": f"{ws} fila {rr}: la fecha {date:%d/%m/%Y} es anterior a la fila previa."})
            prev_date = date
            rows.append(row)
            rr += 1
        blocks.append({"title": title or f"Bloque {a.coordinate}", "consolidated": consolidated, "cell": a.coordinate,
                       "amount": _s(params.get("amount")), "nominal": _s(params.get("nominal")), "units": _s(params.get("units")),
                       "currency": params.get("currency"), "term": _s(params.get("term")), "tnm": _s(params.get("tnm")),
                       "tea": _s(params.get("tea")), "has_status": "status" in cols, "rows": rows})
    return blocks


def _s(v):
    d = D(v)
    return None if d is None else str(d)


def _parse_imc(book, ws):
    sheet = book.wf[ws]
    res = {"sheet": ws, "cut_date": None, "max_rate_pen": None, "max_rate_usd": None, "withholding": None, "holders": [], "overdue": []}
    for row in sheet.iter_rows():
        for c in row:
            t = norm(c.value) if isinstance(c.value, str) else ""
            if t.startswith("fecha de corte"):
                res["cut_date"] = (as_date(book.value(ws, c.row, c.column + 1)) or as_date(book.value(ws, c.row, c.column + 2)) or
                                   as_date(book.value(ws, 1, 1)))
            if t.startswith("tasa maxima"):
                pen, usd = D(book.value(ws, c.row, c.column + 1)), D(book.value(ws, c.row, c.column + 2))
                res["max_rate_pen"] = str((pen * 100).quantize(Decimal("0.0001"))) if pen is not None else None
                res["max_rate_usd"] = str((usd * 100).quantize(Decimal("0.0001"))) if usd is not None else None
            if t == "emision" or t == "emision ":
                hdr = {norm(book.raw(ws, c.row, cc)): cc for cc in range(c.column, c.column + 12) if book.raw(ws, c.row, cc)}
                if "inversionista" in hdr or "bonista" in hdr:
                    name_col = hdr.get("inversionista") or hdr.get("bonista")
                    pct_col = hdr.get("%")
                    rr = c.row + 1
                    while book.raw(ws, rr, c.column):
                        name = str(book.value(ws, rr, name_col) or "").strip()
                        pct = D(book.value(ws, rr, pct_col)) if pct_col else None
                        if name and not any(h["name"] == name and h["series"] == str(book.value(ws, rr, c.column)).strip() for h in res["holders"]):
                            res["holders"].append({"series": str(book.value(ws, rr, c.column)).strip(), "name": name,
                                                   "pct": str(pct) if pct is not None else None})
                        if res["withholding"] is None and hdr.get("retencion ir"):
                            f = str(book.raw(ws, rr, hdr["retencion ir"]) or "")
                            m = re.search(r"(\d+(?:\.\d+)?)%", f)
                            if m:
                                res["withholding"] = m.group(1)
                        rr += 1
                elif "int. moratorio" in hdr:
                    rr = c.row + 1
                    while book.raw(ws, rr, c.column):
                        g = lambda k: book.value(ws, rr, hdr[k]) if k in hdr else None  # noqa: E731
                        res["overdue"].append({"series": str(book.value(ws, rr, c.column)).strip(), "period": _s(g("periodo")),
                                               "date": (as_date(g("fecha")) or datetime.date.min).isoformat(),
                                               "principal": _s(g("principal")), "amort": _s(g("amortizacion")),
                                               "interest": _s(g("intereses")), "moratorio": _s(g("int. moratorio")),
                                               "compensatorio": _s(g("int. compensatorio")), "total": _s(g("total a pagar"))})
                        rr += 1
    if res["cut_date"]:
        res["cut_date"] = res["cut_date"].isoformat()
    return res


def _compare_sheets(schedules):
    issues = []
    if len(schedules) < 2:
        return issues
    by = defaultdict(dict)
    for s in schedules:
        cons = next((b for b in s["blocks"] if b["consolidated"]), None)
        for b in s["blocks"]:
            if b["consolidated"]:
                continue
            for r in b["rows"][1:]:  # la fila 0 es el desembolso
                st = r["status"] if b["has_status"] else _status_from(cons, r["date"])
                by[(series_key(b["title"]), r["label"])][s["sheet"]] = (st or "").lower()
    diffs = defaultdict(list)
    for (serie, label), sts in by.items():
        vals = {v for v in sts.values()}
        if len(sts) > 1 and len(vals) > 1:
            diffs[label].append(f"{serie}: " + ", ".join(f"“{k}” {v or 'sin estado'}" for k, v in sts.items()))
    for label, lst in sorted(diffs.items(), key=lambda kv: kv[0]):
        issues.append({"level": "warning", "msg": f"Cuota {label}: el estado no coincide entre hojas ({'; '.join(lst)})."})
    return issues


def _status_from(cons, date):
    if not cons:
        return None
    return next((r["status"] for r in cons["rows"] if r["date"] == date), None)


def series_key(title):
    m = re.search(r"serie\s+([a-z0-9]+)", norm(title))
    return f"Serie {m.group(1).upper()}" if m else title


# ------------------------------------------------------------------------------- vista previa
def build_preview(parsed, sheet=None, as_of=None):
    """Arma lo que se importaría, con validaciones. No escribe en la base de datos."""
    as_of = as_of or today()
    cfg = SystemSettings.load()
    sheet = sheet or parsed.get("recommended_sheet")
    sched = next((s for s in parsed["schedules"] if s["sheet"] == sheet), None)
    issues = list(parsed.get("issues", []))
    if not sched:
        return {"sheet": sheet, "series": [], "issues": issues + [{"level": "error", "msg": "Seleccione una hoja con cronograma."}]}
    cons = next((b for b in sched["blocks"] if b["consolidated"]), None)
    imc = parsed.get("imc") or {}
    trust_name = parsed.get("trust_name") or "Fideicomiso importado"
    series = []
    for b in sched["blocks"]:
        if b["consolidated"]:
            continue
        s = _series_preview(b, cons, imc, trust_name, as_of, cfg)
        series.append(s)
        issues += s["issues"]
    if cons and series:
        total_amt = sum(Decimal(s["principal"]) for s in series)
        if cons.get("amount") and Decimal(cons["amount"]) != total_amt:
            issues.append({"level": "warning", "msg": f"El monto consolidado ({cons['amount']}) no coincide con la suma de series ({total_amt})."})
    holders_without_series = [h for h in imc.get("holders", []) if not any(series_key(s["title"]) == series_key(h["series"]) for s in series)]
    for h in holders_without_series:
        issues.append({"level": "warning", "msg": f"El bonista {h['name']} pertenece a “{h['series']}”, que no está en la hoja elegida."})
    if not imc.get("holders"):
        issues.append({"level": "review", "msg": "El archivo no trae la distribución por bonista: elija el bonista de cada serie."})
    issues.append({"level": "review", "msg": "El Excel no indica el cliente (fideicomitente): selecciónelo o regístrelo antes de confirmar."})
    if imc.get("holders"):
        issues.append({"level": "review", "msg": "Los bonistas no traen DNI/RUC: se crearán como “Pendiente de registrar” y quedarán en alertas."})
    fee = None
    for sc in [sched] + [x for x in parsed["schedules"] if x is not sched]:  # la comisión puede estar solo en otra hoja
        c = next((b for b in sc["blocks"] if b["consolidated"]), None)
        fees = [Decimal(r["fee"]) for r in (c["rows"] if c else []) if r.get("fee")]
        if fees:
            fee = str(fees[0])
            break
    return {"sheet": sheet, "trust_name": trust_name, "series": series, "issues": _dedupe(issues), "admin_fee": fee,
            "imc": imc, "as_of": as_of.isoformat(), "sheets": [s["sheet"] for s in parsed["schedules"]],
            "recommended": parsed.get("recommended_sheet")}


def _dedupe(issues):
    seen, out = set(), []
    for i in issues:
        k = (i["level"], i["msg"])
        if k not in seen:
            seen.add(k)
            out.append(i)
    return out


def external_ref(trust_name, title):
    return f"{norm(trust_name)}|{norm(series_key(title))}"[:250]


def _series_preview(b, cons, imc, trust_name, as_of, cfg):
    issues = []
    title = b["title"]
    label = series_key(title)
    rows_in = b["rows"]
    if not rows_in:
        return {"title": title, "issues": [{"level": "error", "msg": f"{title}: sin filas."}], "rows": [], "principal": "0"}
    units = int(Decimal(b["units"] or 1))
    tea = Decimal(b["tea"] or 0)
    currency = b["currency"] or "USD"
    if not b["currency"]:
        issues.append({"level": "warning", "msg": f"{title}: moneda no reconocida, se asume USD."})
    start = datetime.date.fromisoformat(rows_in[0]["date"])
    principal = Decimal(rows_in[0]["principal"] or b["amount"] or 0)
    if b["amount"] and principal != Decimal(b["amount"]):
        issues.append({"level": "warning", "msg": f"{title}: el principal inicial ({principal}) difiere del monto de emisión ({b['amount']})."})
    max_rate = Decimal(imc.get("max_rate_usd" if currency == "USD" else "max_rate_pen") or 0) or (cfg.max_rate_usd if currency == "USD" else cfg.max_rate_pen)
    rows = []
    balance = principal
    prev = start
    for i, r in enumerate(rows_in[1:], start=1):
        due = datetime.date.fromisoformat(r["date"])
        days = (due - prev).days
        rate = (ONE + tea) ** (Decimal(days) / Decimal(360)) - ONE
        interest = unit_interest(balance, rate, units)
        amort = money(Decimal(r["amort"] or 0))
        excel_int = Decimal(r["interest_excel"]) if r["interest_excel"] is not None else None
        diff = (excel_int - interest) if excel_int is not None else None
        if diff is not None and abs(diff) > Decimal("0.01"):
            issues.append({"level": "warning", "msg": f"{title} periodo {r['label']}: interés Excel {money(excel_int)} vs. sistema {interest}."})
        status = r["status"] if b["has_status"] else _status_from(cons, r["date"])
        paid = norm(status) == "pagado"
        opening = balance
        balance = money(balance - amort)
        rows.append({"n": i, "label": r["label"] or str(i), "start": prev.isoformat(), "due": due.isoformat(), "days": days,
                     "rate": str(rate.quantize(Decimal("1e-10"))), "opening": str(opening), "amort": str(amort),
                     "interest": str(interest), "interest_excel": str(money(excel_int)) if excel_int is not None else None,
                     "closing": str(balance), "status_excel": status or "", "paid": paid, "row": r["row"]})
        prev = due
    if balance != 0:
        issues.append({"level": "warning", "msg": f"{title}: el saldo final del cronograma no llega a cero ({balance})."})
    # estado automático al día de hoy y penalidad IMC
    for row in rows:
        due = datetime.date.fromisoformat(row["due"])
        cuota = Decimal(row["amort"]) + Decimal(row["interest"])
        if row["paid"]:
            row["state"], row["days_late"], row["penalty"] = "PAGADO", 0, "0.00"
        elif due < as_of:
            res = compute_penalty("IMC" if max_rate else "NONE", max_rate, "INSTALLMENT", 0, 360, due,
                                  {"CAPITAL": Decimal(row["amort"]), "INTEREST": Decimal(row["interest"])}, [], as_of)
            row["state"], row["days_late"], row["penalty"], row["formula"] = "VENCIDO", (as_of - due).days, str(res.amount), res.formula
        else:
            row["state"], row["days_late"], row["penalty"] = "PENDIENTE", 0, "0.00"
        row["cuota"] = str(cuota)
    # verificación contra la hoja IMC (al corte del Excel)
    checks = []
    cut = datetime.date.fromisoformat(imc["cut_date"]) if imc.get("cut_date") else None
    for o in imc.get("overdue", []):
        if series_key(o["series"]) != label or not cut:
            continue
        row = next((r for r in rows if r["due"] == o["date"]), None)
        if not row:
            continue
        res = compute_penalty("IMC", max_rate, "INSTALLMENT", 0, 360, datetime.date.fromisoformat(row["due"]),
                              {"CAPITAL": Decimal(row["amort"]), "INTEREST": Decimal(row["interest"])}, [], cut)
        excel = (Decimal(o["moratorio"] or 0) + Decimal(o["compensatorio"] or 0))
        checks.append({"label": row["label"], "excel": str(money(excel)), "system": str(res.amount), "ok": abs(excel - res.amount) <= Decimal("0.02")})
        if abs(excel - res.amount) > Decimal("0.02"):
            issues.append({"level": "warning", "msg": f"{title} cuota {row['label']}: IMC Excel {money(excel)} vs. sistema {res.amount} al {cut:%d/%m/%Y}."})
    holders = [h for h in imc.get("holders", []) if series_key(h["series"]) == label]
    total_pct = sum((Decimal(h["pct"] or 0) for h in holders), ZERO)
    if holders and abs(total_pct - 1) > Decimal("0.001"):
        issues.append({"level": "warning", "msg": f"{title}: la participación de bonistas suma {total_pct * 100:.2f}% (debería ser 100%)."})
    ref = external_ref(trust_name, title)
    existing = Operation.objects.filter(external_ref=ref).first()
    new_payments = []
    if existing:
        sys_rows = {r.due_date.isoformat(): r for r in existing.current_schedule()}
        for row in rows:
            r = sys_rows.get(row["due"])
            if row["paid"] and r is not None and r.status != "PAGADO":
                new_payments.append(row["label"])
        issues.append({"level": "info", "msg": f"{title} ya existe como {existing.code}: no se duplicará. "
                                               + (f"Se registrarán como pagadas las cuotas {', '.join(new_payments)}." if new_payments
                                                  else "No hay pagos nuevos que registrar.")})
    paid_rows = [r for r in rows if r["paid"]]
    overdue = [r for r in rows if r["state"] == "VENCIDO"]
    upcoming = [r for r in rows if r["state"] == "PENDIENTE"]
    cap_paid = sum((Decimal(r["amort"]) for r in paid_rows), ZERO)
    int_paid = sum((Decimal(r["interest"]) for r in paid_rows), ZERO)
    gen = sum((Decimal(r["interest"]) for r in rows if datetime.date.fromisoformat(r["due"]) <= as_of), ZERO)
    pen = sum((Decimal(r["penalty"]) for r in overdue), ZERO)
    summary = {
        "capital": str(principal), "capital_paid": str(cap_paid), "capital_pending": str(principal - cap_paid),
        "interest_total": str(sum((Decimal(r["interest"]) for r in rows), ZERO)), "interest_generated": str(gen),
        "interest_paid": str(int_paid), "interest_pending": str(max(ZERO, gen - int_paid)), "penalty": str(pen),
        "paid": len(paid_rows), "overdue": len(overdue), "upcoming": len(upcoming),
        "total_due": str(principal - cap_paid + max(ZERO, gen - int_paid) + pen),
        "max_days_late": max((r["days_late"] for r in overdue), default=0),
    }
    return {"title": title, "label": label, "ref": ref, "existing": existing.code if existing else None,
            "existing_id": existing.pk if existing else None, "new_payments": new_payments,
            "currency": currency, "principal": str(principal), "nominal": b["nominal"], "units": units, "term": int(Decimal(b["term"] or len(rows))),
            "tnm": b["tnm"], "tea": str(tea), "start": start.isoformat(), "end": rows[-1]["due"] if rows else None,
            "max_rate": str(max_rate), "rows": rows, "summary": summary, "holders": holders, "checks": checks, "issues": issues}


# ------------------------------------------------------------------------------- confirmación
def file_hash(f):
    h = hashlib.sha256()
    for chunk in f.chunks():
        h.update(chunk)
    f.seek(0)
    return h.hexdigest()


def _match_investor(name):
    key = norm(name)
    for inv in Investor.objects.all():
        if norm(inv.full_name) == key:
            return inv
    return None


def _new_investor(name, user, batch):
    from .portfolio import create_investor
    n = Investor.objects.filter(doc_type="SD").count() + 1
    while Investor.objects.filter(doc_number=f"SD-{n:05d}").exists():
        n += 1
    inv = Investor(first_name=name, doc_type="SD", doc_number=f"SD-{n:05d}", registered_on=timezone.localdate(),
                   notes=f"Creado por la importación {batch.pk} ({batch.original_name}). Falta registrar DNI/RUC y datos bancarios.")
    return create_investor(inv, user)


@transaction.atomic
def confirm_import(batch, user, sheet, client, fallback_investors=None, update_rates=True):
    if batch.status != "PREVIEW":
        raise BusinessError("Esta importación ya fue procesada.")
    preview = build_preview(batch.parsed, sheet)
    if any(i["level"] == "error" for i in preview["issues"]):
        raise BusinessError("Corrija los errores del archivo antes de importar.")
    if client is None and any(not x.get("existing_id") for x in preview["series"]):
        raise BusinessError("Seleccione o registre el cliente (fideicomitente).")
    fallback_investors = fallback_investors or {}
    cfg = SystemSettings.load()
    result = {"created": [], "updated": [], "payments": 0, "investors_created": [], "trust": None, "sheet": sheet}

    if update_rates and preview["imc"]:
        before = snapshot(cfg, ["max_rate_pen", "max_rate_usd", "income_tax_withholding"])
        for field, key in (("max_rate_pen", "max_rate_pen"), ("max_rate_usd", "max_rate_usd"), ("income_tax_withholding", "withholding")):
            v = preview["imc"].get(key)
            if v:
                setattr(cfg, field, Decimal(v))
        cfg.save()
        after = snapshot(cfg, ["max_rate_pen", "max_rate_usd", "income_tax_withholding"])
        if before != after:
            log("UPDATE", "Configuración", cfg, before, after, detail=f"Tasas tomadas de la hoja IMC (importación {batch.pk}).")

    trust_code = "FID-" + re.sub(r"[^A-Z0-9]+", "-", norm(preview["trust_name"]).upper().replace("PATRIMONIO EN FIDEICOMISO", "")).strip("-")[:50]
    trust = Trust.objects.filter(code=trust_code).first()
    if trust is None:
        starts = [datetime.date.fromisoformat(s["start"]) for s in preview["series"]]
        ends = [datetime.date.fromisoformat(s["end"]) for s in preview["series"] if s["end"]]
        trust = Trust.objects.create(
            fiduciary_entity="Por completar", code=trust_code, constitution_date=min(starts), trust_estate=preview["trust_name"],
            settlor=client.full_name if client else "Por completar", trustee="Por completar", beneficiary="Bonistas de la emisión (titulares de valores)",
            contributed_assets="Por completar con el detalle del acto constitutivo.", status="VIGENTE",
            effective_date=min(starts), expiry_date=max(ends) if ends else None, admin_fee=Decimal(preview["admin_fee"] or 0),
            notes=f"Creado por la importación {batch.pk} ({batch.original_name}).", created_by=user)
        log("CREATE", "Fideicomisos", trust, None, snapshot(trust), detail=f"Importación {batch.pk}")
    result["trust"] = trust.code

    for s in preview["series"]:
        if s.get("existing_id"):
            op = Operation.objects.get(pk=s["existing_id"])
            n = _register_excel_payments(op, s, user, batch, only_labels=set(s["new_payments"]))
            if n:
                refresh_operation(op)
                result["updated"].append(op.code)
                result["payments"] += n
            continue
        holders = []
        for h in s["holders"]:
            inv = _match_investor(h["name"])
            if inv is None:
                inv = _new_investor(h["name"], user, batch)
                result["investors_created"].append(inv.full_name)
            holders.append((inv, Decimal(h["pct"] or 0)))
        if not holders:
            inv = fallback_investors.get(s["label"])
            if inv is None:
                raise BusinessError(f"Elija el bonista de {s['title']}.")
            holders = [(inv, Decimal(1))]
        main = max(holders, key=lambda x: x[1])[0]
        max_rate = Decimal(s["max_rate"] or 0)
        rows = s["rows"]
        op = Operation(
            investor=main, client=client, principal=Decimal(s["principal"]), currency=s["currency"],
            disbursement_date=datetime.date.fromisoformat(s["start"]), start_date=datetime.date.fromisoformat(s["start"]),
            maturity_date=datetime.date.fromisoformat(s["end"]), term_months=s["term"], rate_type="TEA",
            rate_value=(Decimal(s["tea"]) * 100).quantize(Decimal("0.0001")), interest_method="COMPOUND", day_base=360,
            day_count="ACTUAL", capitalizations_per_year=12, periodicity="MONTHLY", amortization="BULLET",
            payment_day=None, grace_periods=0, installments_count=len(rows),
            installment_amount=Decimal(next((r["interest"] for r in rows if Decimal(r["amort"]) == 0), rows[0]["interest"])),
            penalty_type="IMC" if max_rate else "NONE", penalty_value=max_rate, penalty_base="INSTALLMENT", penalty_grace_days=0,
            series=s["title"], units=s["units"], nominal_value=Decimal(s["nominal"]) if s["nominal"] else None,
            source="IMPORT", external_ref=s["ref"], import_batch=batch, created_by=user, schedule_version=1,
            purpose=f"{preview['trust_name']} · {s['title']}",
            notes=(f"Importado de “{batch.original_name}”, hoja “{sheet}”. TNM de referencia {Decimal(s['tnm'] or 0) * 100:.2f}%. "
                   "Fechas de vencimiento tomadas del Excel."),
        )
        save_with_code(op, f"{cfg.operation_prefix}-{timezone.localdate().year}-")
        PaymentSchedule.objects.bulk_create([
            PaymentSchedule(operation=op, version=1, number=r["n"], label=r["label"],
                            period_start=datetime.date.fromisoformat(r["start"]), due_date=datetime.date.fromisoformat(r["due"]),
                            days=r["days"], period_rate=Decimal(r["rate"]), opening_balance=Decimal(r["opening"]),
                            capital=Decimal(r["amort"]), interest=Decimal(r["interest"]), closing_balance=Decimal(r["closing"]))
            for r in rows])
        for inv, pct in holders:
            BondHolding.objects.create(operation=op, investor=inv, percentage=pct.quantize(Decimal("0.000001")),
                                       units=int((pct * op.units).quantize(Decimal("1"))))
        trust.operations.add(op)
        log("CREATE", "Operaciones", op, None, snapshot(op),
            detail=f"Importada desde {batch.original_name} (lote {batch.pk}, hoja {sheet}). {len(rows)} cuotas, {len(holders)} bonista(s).")
        result["payments"] += _register_excel_payments(op, s, user, batch)
        refresh_operation(op)
        result["created"].append(op.code)

    batch.status = "CONFIRMED"
    batch.confirmed_by = user
    batch.confirmed_at = timezone.now()
    batch.result = result
    batch.save(update_fields=["status", "confirmed_by", "confirmed_at", "result"])
    log("CREATE", "Importación", batch, None, result, detail=f"Archivo {batch.original_name}")
    return result


def _register_excel_payments(op, s, user, batch, only_labels=None):
    rows = {r.due_date.isoformat(): r for r in op.current_schedule()}
    n = 0
    for r in s["rows"]:
        if not r["paid"] or (only_labels is not None and r["label"] not in only_labels):
            continue
        inst = rows.get(r["due"])
        if inst is None:
            continue
        cap = inst.capital - inst.paid_capital
        intr = inst.interest - inst.paid_interest
        amount = cap + intr
        if amount <= 0:
            continue
        pay = Payment(operation=op, client=op.client, payment_date=inst.due_date, amount=amount, currency=op.currency,
                      concept="IMPORTADO", capital=cap, interest=intr, penalty=ZERO, method="OTRO", registered_by=user,
                      notes=f"Cuota {r['label']} marcada “Pagado” en el Excel (lote {batch.pk}). Fecha real de pago no informada: "
                            "se asume la fecha de vencimiento.")
        save_with_code(pay, f"PAG-{inst.due_date.year}-")
        allocs = []
        if intr > 0:
            allocs.append(PaymentAllocation(payment=pay, installment=inst, component="INTEREST", amount=intr))
        if cap > 0:
            allocs.append(PaymentAllocation(payment=pay, installment=inst, component="CAPITAL", amount=cap))
        PaymentAllocation.objects.bulk_create(allocs)
        log("PAYMENT", "Pagos", pay, None, snapshot(pay), detail=f"{op.code}: importado del Excel")
        n += 1
    return n


def holder_distribution(op, installment, cfg=None):
    """Reparto de una cuota entre bonistas: amortización, interés, IMC, retención de IR y neto."""
    cfg = cfg or SystemSettings.load()
    try:
        imc_amount = installment.penalty.net_amount
    except Exception:  # noqa: BLE001  (cuota sin penalidad)
        imc_amount = ZERO
    wh = cfg.income_tax_withholding / Decimal(100)
    out = []
    holders = list(op.holdings.select_related("investor")) or []
    if not holders:
        holders = [BondHolding(operation=op, investor=op.investor, percentage=Decimal(1), units=op.units)]
    for h in holders:
        p = h.percentage
        amort, intr, imc = money(installment.capital * p), money(installment.interest * p), money(imc_amount * p)
        total = amort + intr + imc
        ret = money((intr + imc) * wh)
        out.append({"holder": h, "pct": p * 100, "amort": amort, "interest": intr, "imc": imc, "total": total,
                    "withholding": ret, "net": total - ret})
    return out
